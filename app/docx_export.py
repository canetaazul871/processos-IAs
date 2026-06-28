"""Exportação da minuta gerada para documento Word (.docx)."""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _is_titulo_secao(linha: str) -> bool:
    """Detecta títulos de seção em caixa-alta (ex.: VOTO, ANÁLISE FIRAC)."""
    if not (3 <= len(linha) <= 50):
        return False
    if not any(c.isalpha() for c in linha):
        return False
    return linha == linha.upper()


def _add_runs_inline(paragrafo, texto: str) -> None:
    """Adiciona texto a um parágrafo convertendo **negrito** em runs em bold."""
    for i, trecho in enumerate(texto.split("**")):
        if not trecho:
            continue
        run = paragrafo.add_run(trecho)
        run.bold = i % 2 == 1  # trechos ímpares estavam entre **...**


def montar_docx(texto: str, titulo: str | None = None) -> bytes:
    """Converte o texto da minuta em um arquivo .docx e devolve os bytes."""
    doc = Document()

    estilo = doc.styles["Normal"]
    estilo.font.name = "Times New Roman"
    estilo.font.size = Pt(12)

    if titulo:
        cab = doc.add_heading(titulo, level=0)
        cab.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for linha in texto.replace("\r\n", "\n").split("\n"):
        bruta = linha.rstrip()
        conteudo = bruta.strip()

        if not conteudo:
            doc.add_paragraph()
            continue
        if set(conteudo) <= {"-", "—", "–", "_", "="} and len(conteudo) >= 3:
            continue  # linha separadora (---)
        if conteudo.startswith("### "):
            doc.add_heading(conteudo[4:].strip(), level=2)
            continue
        if conteudo.startswith("## "):
            doc.add_heading(conteudo[3:].strip(), level=1)
            continue
        if _is_titulo_secao(conteudo):
            doc.add_heading(conteudo, level=1)
            continue

        # Item de lista (- ... ou * ...)
        m = re.match(r"^[-*]\s+(.*)$", conteudo)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_inline(p, m.group(1))
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_runs_inline(p, conteudo)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
